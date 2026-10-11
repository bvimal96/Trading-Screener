import asyncio,time,os
import pandas as pd
from pathlib import Path
from contextlib import asynccontextmanager
from fastapi import FastAPI,Query,Request
from fastapi.responses import HTMLResponse,JSONResponse,StreamingResponse
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi.staticfiles import StaticFiles
from .config import settings
from .universe import build_universe
from .data import YahooData
from .strategies import rolling,weekly,orb
from .alerts import notify


def score_signal(s,intr):
    """Score breakout distance (30), candle strength (20), confirmation (10), and two volume measures (40)."""
    try:
        side=s.get("side")
        entry=float(s.get("entry") or 0)
        ltp=float(s.get("ltp") or 0)
        rg=max(float(s.get("range") or 0),0.01)
        favorable=(ltp-entry) if side=="BUY" else (entry-ltp)
        distance_score=max(0.0,min(30.0,favorable/rg*30.0))
        candle=intr.iloc[-1]
        hi,lo,close=float(candle["High"]),float(candle["Low"]),float(candle["Close"])
        cr=max(hi-lo,0.01)
        close_strength=((close-lo)/cr) if side=="BUY" else ((hi-close)/cr)
        candle_score=max(0.0,min(20.0,close_strength*20.0))
        confirmed=(ltp>entry and close>=entry) if side=="BUY" else (ltp<entry and close<=entry)
        confirm_score=10.0 if confirmed else 0.0
        volumes=intr["Volume"].dropna().astype(float) if "Volume" in intr.columns else []
        avg_vol=float(volumes.mean()) if len(volumes) else 0.0
        current_vol=float(candle.get("Volume",0) or 0)
        relative_volume=current_vol/avg_vol if avg_vol>0 else 0.0
        current_volume_score=max(0.0,min(20.0,relative_volume*10.0))
        breakout_volume_ratio=0.0
        if s.get("entry_time") and avg_vol>0:
            try:
                ts=pd.Timestamp(s["entry_time"])
                if ts.tzinfo is None and getattr(intr.index,"tz",None) is not None:
                    ts=ts.tz_localize(intr.index.tz)
                earlier=intr.loc[intr.index<=ts]
                if not earlier.empty:
                    breakout_volume_ratio=float(earlier.iloc[-1].get("Volume",0) or 0)/avg_vol
            except Exception:
                pass
        breakout_volume_score=max(0.0,min(20.0,breakout_volume_ratio*10.0))
        s["score"]=max(0,min(100,int(round(distance_score+candle_score+confirm_score+current_volume_score+breakout_volume_score))))
        s["breakout_strength"]=round(distance_score+candle_score+confirm_score,1)
        s["relative_volume"]=round(relative_volume,2)
        s["breakout_volume_ratio"]=round(breakout_volume_ratio,2)
    except Exception:
        s["score"]=0
        s["breakout_strength"]=0.0
        s["relative_volume"]=0.0
        s["breakout_volume_ratio"]=0.0
    return s

STATE={"universe":[],"rows":[],"updated":None,"error":None,"previous":set(),"scanning":False}

def calculate(app):
    rows=[]; p=app.state.data
    STATE["scanning"]=True
    try:
        symbols=[u["yahoo"] for u in STATE["universe"]]
        p.prefetch(symbols)
        usable=0
        for u in STATE["universe"]:
            try:
                intr=p.intraday(u["yahoo"])
                d=p.daily(u["yahoo"])
                if intr.empty or d.empty:
                    continue
                usable+=1
                ltp,ts=float(intr.iloc[-1].Close),intr.index[-1]

                ss=[
                    rolling(u["symbol"],u["name"],d,intr,ltp,settings.rolling_rr),
                    weekly(u["symbol"],u["name"],d,intr,ltp,settings.weekly_offset,settings.weekly_rr),
                    orb(u["symbol"],u["name"],d,intr,ltp,settings.weekly_orb_rr,False),
                    orb(u["symbol"],u["name"],d,intr,ltp,settings.monthly_orb_rr,True)
                ]
                for s in ss:
                    if s:
                        score_signal(s,intr)
                        s["updated"]=ts.isoformat() if ts else None
                        rows.append(s)
                        k=f"{s['symbol']}|{s['strategy']}|{s['side']}"
                        if k not in STATE["previous"]:
                            notify(s,os.getenv("PUBLIC_URL",""))
                            STATE["previous"].add(k)
            except Exception:
                continue
        active={f"{x['symbol']}|{x['strategy']}|{x['side']}" for x in rows}
        STATE["previous"] &= active
        # Prioritize index signals at the top: NIFTY 50 first, then BANKNIFTY, then all other symbols.
        priority={"NIFTY":0,"BANKNIFTY":1}
        rows.sort(key=lambda x: (priority.get(x.get("symbol"),2),-x.get("score",0)))
        STATE["rows"]=rows
        STATE["updated"]=time.time()
        print(f"SCAN_OK universe={len(STATE['universe'])} usable={usable} signals={len(rows)}",flush=True)
    finally:
        STATE["scanning"]=False

async def worker(app):
    while True:
        try:
            if not STATE["universe"]:
                STATE["universe"]=await asyncio.to_thread(build_universe)
            await asyncio.to_thread(calculate,app)
            STATE["error"]=None
        except Exception as e:
            STATE["error"]=str(e)
            print(f"SCAN_ERROR {e}",flush=True)
        await asyncio.sleep(settings.poll_seconds)

@asynccontextmanager
async def lifespan(app):
    app.state.data=YahooData(settings.timezone)
    t=asyncio.create_task(worker(app))
    yield
    t.cancel()

app=FastAPI(title="NIFTY Live Breakout Screener",lifespan=lifespan)
app.mount("/static",StaticFiles(directory=Path(__file__).parent.parent/"static"),name="static")

@app.get("/",response_class=HTMLResponse)
def home():
    return (Path(__file__).parent.parent/"static/index.html").read_text()

@app.get("/health")
def health():
    return {"ok":True,"updated":STATE["updated"],"universe":len(STATE["universe"]),
            "signals":len(STATE["rows"]),"scanning":STATE["scanning"],"error":STATE["error"]}

@app.get("/api/status")
def status():
    return {"updated":STATE["updated"],"universe":len(STATE["universe"]),
            "signals":len(STATE["rows"]),"scanning":STATE["scanning"],"error":STATE["error"]}

@app.get("/api/signals")
def signals(strategy:str|None=Query(None),side:str|None=Query(None)):
    r=STATE["rows"]
    if strategy:
        r=[x for x in r if x["strategy"]==strategy]
    if side:
        r=[x for x in r if x["side"]==side]
    priority={"NIFTY":0,"BANKNIFTY":1}
    r=sorted(r,key=lambda x: (priority.get(x.get("symbol"),2),-x.get("score",0)))
    return {"updated":STATE["updated"],"data":r}

@app.post("/api/refresh")
def refresh():
    STATE["universe"]=build_universe()
    calculate(app)
    return {"ok":True,"universe":len(STATE["universe"]),"signals":len(STATE["rows"])}


@app.get("/api/export.xlsx")
def export_xlsx(
    strategy:str|None=Query(None),
    side:str|None=Query(None),
    status:str|None=Query(None),
    search:str|None=Query(None)
):
    rows=list(STATE["rows"])
    if strategy:
        rows=[x for x in rows if x.get("strategy")==strategy]
    if side:
        rows=[x for x in rows if x.get("side")==side]
    if status:
        rows=[x for x in rows if x.get("status")==status]
    if search:
        q=search.strip().upper()
        rows=[x for x in rows if q in str(x.get("symbol","")).upper() or q in str(x.get("name","")).upper()]
    priority={"NIFTY":0,"BANKNIFTY":1}
    rows.sort(key=lambda x: (priority.get(x.get("symbol"),2),-x.get("score",0)))

    columns=[
        ("Symbol","symbol"),("Company Name","name"),("Strategy","strategy"),
        ("Side","side"),("Score","score"),("Breakout Strength","breakout_strength"),("Relative Volume","relative_volume"),("Breakout Volume Ratio","breakout_volume_ratio"),("Entry Price","entry"),("LTP","ltp"),("P&L (Points)","pnl"),
        ("Stop Loss","sl"),("Target","target"),("Range","range"),
        ("Status","status"),
        ("Entry Time","entry_time"),("Event Time","event_time"),("Last Updated","updated")
    ]
    strategy_labels={
        "Rolling 2-Day":"Roll Intra Day",
        "Weekly High/Low":"Week HL- EOW",
        "Weekly ORB":"Week 2 Day",
        "Monthly ORB":"Month 2 Day"
    }
    wb=Workbook()
    ws=wb.active
    ws.title="Screener Signals"
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="1F4E78")
        cell.alignment=Alignment(horizontal="center")
    for row in rows:
        values=[]
        for _,key in columns:
            if key=="pnl":
                try:
                    entry=float(row.get("entry")); ltp=float(row.get("ltp"))
                    value=round((ltp-entry) if row.get("side")=="BUY" else (entry-ltp),2)
                except (TypeError,ValueError):
                    value=None
            elif key=="strategy":
                value=strategy_labels.get(row.get(key),row.get(key))
            else:
                value=row.get(key)
            values.append(value)
        ws.append(values)
    ws.freeze_panes="A2"
    ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        letter=col[0].column_letter
        max_len=max((len(str(cell.value or "")) for cell in col),default=10)
        ws.column_dimensions[letter].width=min(max(max_len+2,12),32)
    ws.sheet_view.showGridLines=False
    buffer=BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename="screener_signals_"+datetime.now().strftime("%Y%m%d_%H%M%S")+".xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":f'attachment; filename="{filename}"'}
    )

@app.post("/webhooks/whatsapp")
async def whatsapp_webhook(request:Request):
    return JSONResponse({"ok":True})
